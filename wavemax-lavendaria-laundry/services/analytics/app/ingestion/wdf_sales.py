from __future__ import annotations

import json
import math
from collections.abc import Iterator, Sequence
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from hashlib import sha256
from pathlib import Path

from openpyxl import load_workbook

from app.ingestion.models import OutcomeStatus, RawCell, RawRecord, RowOutcome, WdfSalesLine

PARSER_VERSION = "wdf-sales-xlsx/1.0.0"
RULE_VERSION = "wdf-sales-quality/1.0.0"
SHEET_NAME = "Sales Breakdown"
EXPECTED_HEADERS = (
    "Paid Date",
    "Order #",
    "Store",
    "Category",
    "Department",
    "Line Item Name",
    "Quantity",
    "Units",
    "Line Item Subtotal",
    "Tax 1",
    "Tax 2",
    "Tax 3",
    "Line Item Note",
    "Laundry Service Item",
)


class SourceSchemaError(ValueError):
    """Raised when a workbook cannot be safely identified as the approved source kind."""


def parse_wdf_sales_workbook(path: Path, *, source_sha256: str) -> Iterator[RowOutcome]:
    if path.suffix.lower() != ".xlsx":
        raise SourceSchemaError("WDF annual sales source must use the .xlsx format")

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if SHEET_NAME not in workbook.sheetnames:
            raise SourceSchemaError(f"Required worksheet is missing: {SHEET_NAME}")
        sheet = workbook[SHEET_NAME]
        rows = sheet.iter_rows(values_only=True)
        try:
            raw_headers = next(rows)
        except StopIteration as error:
            raise SourceSchemaError("Sales Breakdown worksheet is empty") from error

        headers = tuple(_optional_text(value) for value in raw_headers)
        if headers != EXPECTED_HEADERS:
            raise SourceSchemaError("Sales Breakdown headers do not match the approved schema")

        seen_record_hashes: set[str] = set()
        for source_row_number, values in enumerate(rows, start=2):
            raw_record = _fixed_width_record(values)
            raw_record_hash = _hash_record(raw_record)
            warnings: list[str] = []
            if raw_record_hash in seen_record_hashes:
                warnings.append("EXACT_CONTENT_REPEAT")
            else:
                seen_record_hashes.add(raw_record_hash)

            parsed, reasons = _parse_row(raw_record)
            identity = sha256(
                f"{source_sha256}:{SHEET_NAME}:{source_row_number}".encode()
            ).hexdigest()
            status = OutcomeStatus.QUARANTINED if reasons else OutcomeStatus.ACCEPTED
            yield RowOutcome(
                source_row_number=source_row_number,
                source_row_identity=identity,
                raw_record_hash=raw_record_hash,
                raw_record=raw_record,
                status=status,
                reason_codes=tuple(sorted(set(reasons))),
                warning_codes=tuple(warnings),
                parsed=parsed if status is OutcomeStatus.ACCEPTED else None,
            )
    finally:
        workbook.close()


def _fixed_width_record(values: Sequence[RawCell]) -> RawRecord:
    width = len(EXPECTED_HEADERS)
    padded = tuple(values[:width]) + (None,) * max(0, width - len(values))
    return padded[:width]


def _parse_row(raw: RawRecord) -> tuple[WdfSalesLine | None, list[str]]:
    if all(_is_blank(value) for value in raw):
        return None, ["ROW_EMPTY"]

    reasons: list[str] = []
    paid_date = _date_value(raw[0], "PAID_DATE", reasons)
    order_number = _required_identifier(raw[1], "ORDER_NUMBER", reasons)
    store = _required_text(raw[2], "STORE", reasons)
    line_item_name = _required_text(raw[5], "LINE_ITEM_NAME", reasons)
    quantity = _decimal_value(raw[6], "QUANTITY", reasons)
    subtotal = _decimal_value(raw[8], "LINE_ITEM_SUBTOTAL", reasons)
    tax_1 = _decimal_value(raw[9], "TAX_1", reasons)
    tax_2 = _decimal_value(raw[10], "TAX_2", reasons)
    tax_3 = _decimal_value(raw[11], "TAX_3", reasons)
    laundry_service_item = _boolean_value(raw[13], "LAUNDRY_SERVICE_ITEM", reasons)

    if reasons:
        return None, reasons
    assert paid_date is not None
    assert order_number is not None
    assert store is not None
    assert line_item_name is not None
    assert quantity is not None
    assert subtotal is not None
    assert tax_1 is not None
    assert tax_2 is not None
    assert tax_3 is not None
    assert laundry_service_item is not None
    return (
        WdfSalesLine(
            paid_date=paid_date,
            order_number=order_number,
            store=store,
            category=_optional_text(raw[3]),
            department=_optional_text(raw[4]),
            line_item_name=line_item_name,
            quantity=quantity,
            units=_optional_text(raw[7]),
            line_item_subtotal=subtotal,
            tax_1=tax_1,
            tax_2=tax_2,
            tax_3=tax_3,
            line_item_note=_optional_text(raw[12]),
            laundry_service_item=laundry_service_item,
        ),
        [],
    )


def _is_blank(value: RawCell) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def _optional_text(value: RawCell) -> str | None:
    if _is_blank(value):
        return None
    return str(value).strip()


def _required_text(value: RawCell, field: str, reasons: list[str]) -> str | None:
    result = _optional_text(value)
    if result is None:
        reasons.append(f"{field}_MISSING")
    return result


def _required_identifier(value: RawCell, field: str, reasons: list[str]) -> str | None:
    if isinstance(value, bool) or _is_blank(value):
        reasons.append(f"{field}_MISSING")
        return None
    if isinstance(value, float):
        if not math.isfinite(value):
            reasons.append(f"{field}_INVALID")
            return None
        if value.is_integer():
            return str(int(value))
    return str(value).strip()


def _date_value(value: RawCell, field: str, reasons: list[str]) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    reasons.append(f"{field}_{'MISSING' if _is_blank(value) else 'INVALID'}")
    return None


def _decimal_value(value: RawCell, field: str, reasons: list[str]) -> Decimal | None:
    if isinstance(value, bool) or _is_blank(value):
        reasons.append(f"{field}_{'MISSING' if _is_blank(value) else 'INVALID'}")
        return None
    try:
        result = Decimal(str(value))
    except (InvalidOperation, ValueError):
        reasons.append(f"{field}_INVALID")
        return None
    if not result.is_finite():
        reasons.append(f"{field}_INVALID")
        return None
    return result


def _boolean_value(value: RawCell, field: str, reasons: list[str]) -> bool | None:
    if isinstance(value, bool):
        return value
    reasons.append(f"{field}_{'MISSING' if _is_blank(value) else 'INVALID'}")
    return None


def _hash_record(raw_record: RawRecord) -> str:
    encoded = json.dumps(
        [_json_value(value) for value in raw_record],
        separators=(",", ":"),
        ensure_ascii=True,
    ).encode()
    return sha256(encoded).hexdigest()


def _json_value(value: RawCell) -> str | int | float | bool | None:
    if isinstance(value, datetime | date):
        return value.isoformat()
    if isinstance(value, Decimal):
        return str(value)
    return value
